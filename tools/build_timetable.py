# 教科担任別時間割（Excel）から、時間割ページ用のデータ site/jikanwari_data.js を作る。
# 使い方：python3 tools/build_timetable.py [Excelファイル]
# 必要なもの：pip install openpyxl
#
# 読み取る Excel の形（source/時間割_見本.xlsx を参照）
#   シート「教科担任別一覧」：1行に先生1人。B列＝担任の組（1A など。IJ は特別支援学級）、C列＝名前、
#     D列から 月①〜⑥、1列あけて 火①〜⑥ … のように、曜日ごとに 7 列ずつ。セルには「1A」「2CD」（合同）、
#     「1CT」（TT・メイン）「1Ct」（TT・援助）、「IJ数」（IJ の教科）や「主」「道」などの記号を書く。
#   シート「クラス別一覧（入力）」：1行に1クラス。B列＝組、C列から 月①〜⑥ … の教科（国・数 など）。
# 学校によって形が違うときは、下の「学校ごとの設定」を直すか、Claude に合わせてもらう。
import hashlib, json, re, sys, unicodedata
from collections import Counter
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
SRC = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "source/時間割_見本.xlsx"

# ===== 学校ごとの設定 =====
TEACHER_SHEET, CLASS_SHEET = "教科担任別一覧", "クラス別一覧（入力）"
TEACHER_FIRST_ROW, CLASS_FIRST_ROW = 7, 9   # データが始まる行
VERSION_CELL = "AD3"                        # 版の日付が入っているセル（無ければそのままでよい）
OUT = ROOT / "site/jikanwari_data.js"

DAYS = "月火水木金"
# 教科担任別一覧：月①が D 列（4列目）、曜日ごとに 7 列ずつ（6時限＋区切り1列）
TEACHER_COL = lambda d, p: 4 + d * 7 + p
# クラス別一覧（入力）：月①が C 列（3列目）
CLASS_COL = lambda d, p: 3 + d * 7 + p

SUBJECT = {"国": "国語", "国Ⅰ": "国語Ⅰ", "国Ⅱ": "国語Ⅱ", "社": "社会", "数": "数学", "理": "理科", "英": "英語",
           "音": "音楽", "美": "美術", "保体": "保体", "技": "技術", "家": "家庭", "技家": "技術・家庭",
           "音美": "音楽・美術", "美音": "音楽・美術", "自": "自立活動", "作": "作業学習",
           "道": "道徳", "学": "学活", "総": "総合"}
# 時間割の記号 → カレンダーに出す名前（学校に合わせて直す）
SPECIAL = {"主": "主任会", "生": "生活係会", "待機": "待機", "初": "初任者研修", "さくら": "さくら",
           "道": "道徳", "学": "学活", "総": "総合"}
# IJ（特別支援学級）の時間として扱う記号
IJ_SPECIAL = {"T2": "IJ・T2"}
SPLIT = {"音美": ("音", "美"), "美音": ("音", "美"), "技家": ("技", "家")}


def norm(v):
    if v is None or v == 0:
        return ""
    return unicodedata.normalize("NFKC", str(v)).strip()


wb = openpyxl.load_workbook(SRC, data_only=True)
ws_t, ws_c = wb[TEACHER_SHEET], wb[CLASS_SHEET]
version = ws_t[VERSION_CELL].value
version = f"{version.month}/{version.day}版" if hasattr(version, "month") else ""

# クラスごとの教科（例：CLASS["1A"][0][0] = 月①の教科）
CLASS = {}
for r in range(CLASS_FIRST_ROW, ws_c.max_row + 1):
    name = norm(ws_c.cell(r, 2).value).replace("・", "")
    if not name:
        continue
    CLASS[name] = [[norm(ws_c.cell(r, CLASS_COL(d, p)).value) for p in range(6)] for d in range(5)]

problems = []
teachers = []
for r in range(TEACHER_FIRST_ROW, ws_t.max_row + 1):
    name = norm(ws_t.cell(r, 3).value)
    raw = [[norm(ws_t.cell(r, TEACHER_COL(d, p)).value) for p in range(6)] for d in range(5)]
    if not name or not any(any(x) for x in raw):
        continue
    home = norm(ws_t.cell(r, 2).value)

    # 学年の判定と、音美・技家のどちらを教えているか（他のコマの教科から多数決）
    grades, prefer = Counter(), Counter()
    for d in range(5):
        for p in range(6):
            m = re.fullmatch(r"([123])([A-F]+)([Tt]?)", raw[d][p])
            if m:
                grades[m.group(1)] += 1
                for c in m.group(2):
                    s = CLASS.get(m.group(1) + c, [[""] * 6] * 5)[d][p]
                    if s and s not in SPLIT:
                        prefer[s] += 1
    if home[:1] in "123" and home:
        base = home[0]
    elif home == "IJ":
        base = "IJ"
    else:
        base = grades.most_common(1)[0][0] if grades else "1"

    def subject_of(cls, d, p):
        s = CLASS.get(cls, [[""] * 6] * 5)[d][p]
        if s in SPLIT:
            a, b = SPLIT[s]
            if prefer[a] or prefer[b]:
                s = a if prefer[a] >= prefer[b] else b
        return SUBJECT.get(s, s)

    slots = []
    for d in range(5):
        row = []
        for p in range(6):
            t = raw[d][p]
            where = f"{name} {DAYS[d]}{p + 1}「{t}」"
            if not t:
                row.append(None)
            elif t in SPECIAL:
                row.append([base, SPECIAL[t]])
            elif t in IJ_SPECIAL:
                row.append(["IJ", IJ_SPECIAL[t]])
            elif t == "IJ":
                s = subject_of("IJ", d, p)
                row.append(["IJ", f"IJ {s}"])
                if not s:
                    problems.append(where + "：IJ の教科が見つかりません")
            elif t.startswith("IJ"):
                s = t[2:]
                row.append(["IJ", f"IJ {SUBJECT.get(s, s)}"])
            else:
                m = re.fullmatch(r"([123])([A-F]+)([Tt]?)", t)
                if not m:
                    problems.append(where + "：読めない記号（そのまま登録）")
                    row.append([base, t])
                    continue
                g, letters, tt = m.groups()
                classes = [g + c for c in sorted(letters)]
                subs = []
                for c in classes:
                    s = subject_of(c, d, p)
                    if not s:
                        problems.append(where + f"：{c} の教科が空欄")
                    if s and s not in subs:
                        subs.append(s)
                label = "・".join(classes) + " " + "・".join(subs)
                if tt:
                    label += "（TT）" if tt == "T" else "（TT援助）"
                row.append([g, label.strip()])
        slots.append(row)
    tid = hashlib.sha1(name.encode()).hexdigest()[:8]
    teachers.append({"id": tid, "name": name, "home": home if home[:1] in "123" or home == "IJ" else "", "base": base, "slots": slots})

js = ("// 教科担任別時間割から自動で作ったデータ。手で直さず tools/build_timetable.py を実行し直すこと。\n"
      f"// 元：{SRC.name}（{version}）\n"
      f"var TIMETABLE_VERSION = {json.dumps(version, ensure_ascii=False)};\n"
      "var TEACHERS = [\n" + ",\n".join(json.dumps(t, ensure_ascii=False, separators=(",", ":")) for t in teachers) + "\n];\n"
      "if (typeof module !== \"undefined\") module.exports = { TEACHERS, TIMETABLE_VERSION };\n")
OUT.write_text(js, encoding="utf-8")
print(f"{OUT.relative_to(ROOT)}：{len(teachers)}人（{version}）")
for p in problems:
    print("  要確認：" + p)
